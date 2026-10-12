import json, runpy, re
from pathlib import Path
from html.parser import HTMLParser
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, Image, Spacer, PageBreak, KeepTogether
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\decision\work\pdf-report')
captured=[]
def encode(flow):
    if isinstance(flow, Paragraph):return {'kind':'paragraph','text':flow.text,'style':flow.style.name}
    if isinstance(flow, PageBreak):return {'kind':'pagebreak'}
    if isinstance(flow, Image):return {'kind':'image','path':flow.filename,'width':flow.drawWidth,'height':flow.drawHeight}
    if isinstance(flow, KeepTogether):return {'kind':'group','items':[encode(x) for x in flow._content]}
    if isinstance(flow, Spacer):return {'kind':'spacer','height':flow.height}
    if isinstance(flow, Table):
        return {'kind':'table','widths':flow._colWidths,'cells':[[[encode(v) for v in cell] if isinstance(cell,(list,tuple)) else [encode(cell)] for cell in row] for row in flow._cellvalues],
                'backgrounds':[[c[0],list(c[1]),list(c[2]),c[3].hexval() if hasattr(c[3],'hexval') else None] for c in flow._bkgrndcmds if c[0]=='BACKGROUND']}
    raise TypeError(type(flow))
real_build=SimpleDocTemplate.build
def capture(self,flowables,*a,**kw):
    captured.extend(encode(x) for x in flowables)
    return real_build(self,flowables,*a,**kw)
SimpleDocTemplate.build=capture
runpy.run_path(str(ROOT/'build_threeway.py'),run_name='__main__')
SimpleDocTemplate.build=real_build
(ROOT/'threeway/report_content.json').write_text(json.dumps(captured,ensure_ascii=False,indent=2),encoding='utf-8')

doc=Document();sec=doc.sections[0]
sec.page_width=Pt(595.276);sec.page_height=Pt(841.89)
sec.left_margin=sec.right_margin=Pt(44);sec.top_margin=Pt(65);sec.bottom_margin=Pt(51)
sec.header_distance=Pt(30);sec.footer_distance=Pt(23)
def setfont(font,size,bold=False):
    font.name='Microsoft YaHei';font.size=Pt(size);font.bold=bold;font.color.rgb=RGBColor.from_string('000000')
for name,size,bold in [('Normal',9.6,False),('Title',25,True),('Heading 1',18,True),('Heading 2',11.5,True),('Caption',8,False)]:
    st=doc.styles[name];setfont(st.font,size,bold)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=Pt(15.8 if name=='Normal' else 34 if name=='Title' else 25 if name=='Heading 1' else 17 if name=='Heading 2' else 12)
    st.paragraph_format.keep_with_next=name.startswith('Heading')
    st.paragraph_format.widow_control=True
doc.styles['Normal'].paragraph_format.space_before=Pt(0)
for name in ['Title','Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(0)
header=sec.header.paragraphs[0];header.text='决策模型与基础模型三方评估报告';setfont(header.runs[0].font,8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for text in ['第 ',None,' 页']:
    if text is not None:setfont(footer.add_run(text).font,8)
    else:
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='决策模型与基础模型准确率和响应速度对比'
doc.core_properties.author='Decision Benchmark'
doc.core_properties.subject='阿里云百炼 Microsoft Decision-1 Qwen3.5-9B 三方实测'

class RichText(HTMLParser):
    def __init__(self,p,size,bold=False):
        super().__init__(convert_charrefs=True);self.p=p;self.size=size;self.base=bold;self.depth=0
    def handle_starttag(self,tag,attrs):
        if tag in ('b','strong'):self.depth+=1
        elif tag=='br':self.p.add_run().add_break()
    def handle_endtag(self,tag):
        if tag in ('b','strong'):self.depth=max(0,self.depth-1)
    def handle_data(self,data):
        r=self.p.add_run(data);setfont(r.font,self.size,self.base or self.depth>0)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')

def paragraph(item,container=doc,existing=None,cell=False):
    name=item['style'];text=item['text']
    style={'TitleCN':'Title','SectionCN':'Heading 1','SubCN':'Heading 2','SmallCN':'Caption'}.get(name,'Normal')
    size={'TitleCN':25,'SectionCN':18,'SubCN':11.5,'SmallCN':8,'CodeCN':8.2,'Cell':8.2,'CellBold':8.2}.get(name,9.6)
    p=existing if existing is not None else container.add_paragraph()
    p.style=style;p.paragraph_format.keep_with_next=name in ['TitleCN','SectionCN','SubCN']
    if name=='SectionCN':text=text.replace('  /  ','  ').replace('：',' ').replace('，',' ')
    if name=='TitleCN':text=text.replace('、','和')
    if name=='SubCN':text=text.replace('：',' ').replace('，',' ').replace(' / ',' ')
    p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.space_after=Pt(0 if cell else 5 if name=='SmallCN' else 6 if name=='CodeCN' else 7)
    if cell or name in ['Cell','CellBold']:p.paragraph_format.line_spacing=Pt(12.3)
    if name=='CodeCN':p.paragraph_format.line_spacing=Pt(13)
    RichText(p,size,name in ['TitleCN','SectionCN','SubCN','CellBold']).feed(text)
    return p

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr();el=OxmlElement('w:shd');el.set(qn('w:fill'),fill.replace('0x','').replace('#','').upper());pr.append(el)

def render_table(item):
    rows=item['cells'];widths=item['widths']
    if len(rows)==1 and len(rows[0])==1:
        # Summary and audit text stays editable prose, rather than a nested layout table.
        entries=rows[0][0]
        ps=[]
        for entry in entries:
            if entry['kind']=='paragraph':ps.append(paragraph(entry))
        for p in ps[:-1]:p.paragraph_format.keep_with_next=True
        if ps:ps[-1].paragraph_format.space_after=Pt(9)
        return
    table=doc.add_table(rows=len(rows),cols=len(widths));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    pr=table._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    for col,w in zip(table.columns,widths):col.width=Pt(w)
    for i,row in enumerate(table.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        if i==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for j,cell in enumerate(row.cells):
            cell.width=Pt(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            marg=OxmlElement('w:tcMar')
            for side,value in [('top',95),('bottom',95),('left',130),('right',130)]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');marg.append(el)
            cell._tc.get_or_add_tcPr().append(marg)
            shade(cell,'E7EEF5' if i==0 else 'F8FAFC' if i%2==0 else 'FFFFFF')
            for k,part in enumerate(rows[i][j]):
                if part['kind']=='paragraph':paragraph(part,cell,cell.paragraphs[0] if k==0 else None,True)
    for typ,start,end,color in item['backgrounds']:
        if start==end and start[1]>0 and color:
            cell=table.cell(start[1],start[0]);shade(cell,color)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(1);setfont(p.add_run().font,1)

def render(items):
    for item in items:
        kind=item['kind']
        if kind=='paragraph':paragraph(item)
        elif kind=='table':render_table(item)
        elif kind=='group':render(item['items'])
        elif kind=='pagebreak':
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(1)
            p.add_run().add_break(__import__('docx.enum.text',fromlist=['WD_BREAK']).WD_BREAK.PAGE)
        elif kind=='image':
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=True
            run=p.add_run();run.font.size=Pt(1);pic=run.add_picture(item['path'],width=Pt(item['width']),height=Pt(item['height']))
            pic._inline.docPr.set('descr',Path(item['path']).stem+' 三方测试数据图示')
        elif kind!='spacer':raise ValueError(kind)
render(captured)
out=Path(r'C:\decision\reports\decision_models_threeway_2026-10-11.docx')
for tree in [doc.styles.element,doc.element]:
    for border in list(tree.xpath('.//w:pBdr')):
        border.getparent().remove(border)
doc.save(out)
from zipfile import ZipFile
with ZipFile(out) as z:
    xml=z.read('word/document.xml').decode()
    assert not any(t in xml for t in ['北京直连','北京直連','直连','直連'])
    assert 'DCF2DC' in xml
print(json.dumps({'docx':str(out),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes)},ensure_ascii=False))
